"""Chapter 12: Back to Excel

Code from The Excel-to-Python Bridge by Mathew Kariuki.
Run from the companion folder, or paste into a notebook cell by cell.
"""

# --- Step 1 ---
import pandas as pd
from pathlib import Path

xl = "data/kestrel_sales.xlsx"
orders = pd.read_excel(xl, sheet_name="Orders")
customers = pd.read_excel(xl, sheet_name="Customers")
orders["Revenue"] = (orders["Qty"] * orders["UnitPrice"]
                     * (1 - orders["Discount"])).round(2)
orders = orders.merge(customers[["CustomerID", "State"]],
                      on="CustomerID", how="left")

out = Path("output")
out.mkdir(exist_ok=True)

# --- Step 2 ---
by_rep = (orders.groupby("Rep", as_index=False)
          .agg(Orders=("OrderID", "nunique"),
               Revenue=("Revenue", "sum")))
by_rep.to_excel(out / "rep_summary.xlsx", index=False)
pd.read_excel(out / "rep_summary.xlsx")

# --- Step 3 ---
by_state = (orders.groupby("State", as_index=False)
            ["Revenue"].sum())
with pd.ExcelWriter(out / "summary.xlsx") as writer:
    by_rep.to_excel(writer, sheet_name="By rep", index=False)
    by_state.to_excel(writer, sheet_name="By state",
                      index=False)

pd.ExcelFile(out / "summary.xlsx").sheet_names

# --- Step 4 ---
with pd.ExcelWriter(out / "rep_report.xlsx",
                    engine="xlsxwriter") as writer:
    by_rep.to_excel(writer, sheet_name="Reps", index=False,
                    startrow=1, header=False)
    wb = writer.book
    ws = writer.sheets["Reps"]

    head = wb.add_format({"bold": True, "bg_color": "#D9D9D9",
                          "bottom": 1})
    money = wb.add_format({"num_format": "$#,##0"})
    count = wb.add_format({"num_format": "#,##0"})

    for col, name in enumerate(by_rep.columns):
        ws.write(0, col, name, head)
    ws.set_column("A:A", 16)
    ws.set_column("B:B", 10, count)
    ws.set_column("C:C", 14, money)
    ws.freeze_panes(1, 0)
    ws.autofilter(0, 0, len(by_rep), len(by_rep.columns) - 1)
print("saved")

# --- Step 5 ---
pt = orders.pivot_table(index="State", columns="Rep",
                        values="Revenue", aggfunc="sum",
                        fill_value=0).round(0)
with pd.ExcelWriter(out / "heatmap.xlsx",
                    engine="xlsxwriter") as writer:
    pt.to_excel(writer, sheet_name="Heatmap")
    ws = writer.sheets["Heatmap"]
    n_rows, n_cols = pt.shape
    ws.conditional_format(1, 1, n_rows, n_cols,
                          {"type": "3_color_scale"})
    ws.set_column(0, n_cols, 14)
print(pt.shape)

# --- Step 6 ---
with pd.ExcelWriter(out / "with_total.xlsx",
                    engine="xlsxwriter") as writer:
    by_rep.to_excel(writer, sheet_name="Reps", index=False)
    ws = writer.sheets["Reps"]
    last = len(by_rep) + 1          # last data row, Excel-style
    ws.write(last, 0, "Total")
    ws.write_formula(last, 2, f"=SUM(C2:C{last})")

from openpyxl import load_workbook
ws = load_workbook(out / "with_total.xlsx")["Reps"]
ws["A7"].value, ws["C7"].value

# --- Step 7 ---
with pd.ExcelWriter(out / "summary.xlsx", mode="a",
                    engine="openpyxl",
                    if_sheet_exists="replace") as writer:
    by_state.sort_values("Revenue", ascending=False).to_excel(
        writer, sheet_name="By state", index=False)

pd.ExcelFile(out / "summary.xlsx").sheet_names

